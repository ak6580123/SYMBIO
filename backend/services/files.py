"""Minimal report text extraction for PDF / TXT / DOCX uploads.

Deliberately small: no OCR, no layout analysis, no third-party conversion
services. If extraction fails the API returns ok=false and the UI asks the user
to type the description instead - the flow never dead-ends.

PDF engines, in order of preference:

  1. PyMuPDF (`fitz`)   - better text extraction on real-world reports, but a
                          ~25 MB binary wheel, so it is OPTIONAL.
  2. pypdf              - pure Python, already required, always available.

Install PyMuPDF only if your reports extract poorly:  pip install PyMuPDF
"""

from __future__ import annotations

import io

MAX_BYTES = 8 * 1024 * 1024
MAX_CHARS = 12000
SUPPORTED = ("pdf", "txt", "docx")


def _pymupdf():
    """Import PyMuPDF if installed. `pymupdf` is the modern module name;
    `fitz` is kept for installs pinned to an older release."""
    try:
        import pymupdf

        return pymupdf
    except ImportError:
        import fitz

        return fitz


def pdf_engine() -> str:
    """Name of the PDF engine in use - surfaced in the upload response."""
    try:
        _pymupdf()
        return "pymupdf"
    except ImportError:
        return "pypdf"


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract(filename: str, raw: bytes) -> dict:
    """Return {filename, extracted_text, char_count, ok, message}."""
    kind = _ext(filename)

    if len(raw) > MAX_BYTES:
        return {
            "filename": filename,
            "extracted_text": "",
            "char_count": 0,
            "ok": False,
            "message": f"File is larger than {MAX_BYTES // (1024 * 1024)} MB. Please upload a smaller report.",
        }

    if kind not in SUPPORTED:
        return {
            "filename": filename,
            "extracted_text": "",
            "char_count": 0,
            "ok": False,
            "message": f"Unsupported file type '.{kind or 'unknown'}'. Symbio accepts PDF, TXT and DOCX.",
        }

    try:
        if kind == "pdf":
            text = _from_pdf(raw)
        elif kind == "docx":
            text = _from_docx(raw)
        else:
            text = raw.decode("utf-8", errors="ignore")
    except Exception as exc:  # noqa: BLE001 - any parse failure must degrade gracefully
        return {
            "filename": filename,
            "extracted_text": "",
            "char_count": 0,
            "ok": False,
            "message": f"Could not read the file ({exc.__class__.__name__}). Please type the description instead.",
        }

    text = " ".join(text.split())[:MAX_CHARS]
    if not text:
        return {
            "filename": filename,
            "extracted_text": "",
            "char_count": 0,
            "ok": False,
            "message": "No selectable text found - this looks like a scanned image. Please type the description instead.",
        }

    return {
        "filename": filename,
        "extracted_text": text,
        "char_count": len(text),
        "ok": True,
        "message": f"Extracted {len(text):,} characters from {filename}.",
        "engine": pdf_engine() if kind == "pdf" else None,
    }


def _from_pdf(raw: bytes) -> str:
    try:
        return _from_pdf_pymupdf(raw)
    except ImportError:
        return _from_pdf_pypdf(raw)


def _from_pdf_pymupdf(raw: bytes) -> str:
    pymupdf = _pymupdf()  # optional dependency

    pages = []
    with pymupdf.open(stream=raw, filetype="pdf") as document:
        for page in document[:40]:
            pages.append(page.get_text())
    return "\n".join(pages)


def _from_pdf_pypdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = [(page.extract_text() or "") for page in reader.pages[:40]]
    return "\n".join(pages)


def _from_docx(raw: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables[:10]:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
